"""Regression tests for fd_reader.report using entirely synthetic
GuestRecord/ReservationRecord data -- no real guest-notes/ PDFs involved.
Covers the same structural cases previously pinned against real sample
data (blue placeholder for guests with nothing, red for hard mismatches,
yellow for room/party-size disagreements, linked-group and room-move
notes, perk badges, xlsx output) but with made-up guests (see
tests_synthetic/fixtures.py).
"""
from datetime import datetime

import openpyxl

from fd_reader.match import cross_check_notes, detect_room_moves, find_room_groups, match_reservations
from fd_reader.report import (
    BLOCK_BORDER_FILL,
    BLUE,
    GREEN,
    ORANGE,
    RED,
    YELLOW,
    build_guest_blocks,
    write_report,
)
from fd_reader.report.xlsx_writer import LEFT_COL_COUNT

from tests_synthetic.fixtures import make_guest, make_reservation


def by_name(blocks, name_substring):
    matches = [b for b in blocks if name_substring.lower() in b.guest.guest_name.lower()]
    assert matches, f"no block found for {name_substring!r}"
    return matches


def build(guests, reservations):
    match_results = match_reservations(guests, reservations)
    note_checks = cross_check_notes(guests, match_results)
    room_moves = detect_room_moves(guests)
    return build_guest_blocks(guests, match_results, note_checks, room_moves), room_moves


def test_every_guest_gets_exactly_one_block():
    guests = [
        make_guest(confirmation_number="A-1", guest_name="Jackson, John"),
        make_guest(confirmation_number="A-2", guest_name="Sheeran, Steve", room_name="BUCK"),
    ]
    blocks, _ = build(guests, [])
    assert len(blocks) == 2


def test_blocks_sorted_by_arrival_date():
    guests = [
        make_guest(confirmation_number="B-1", guest_name="Depp, Emma", arrival_date="07/15/2026"),
        make_guest(confirmation_number="B-2", guest_name="Hanks, Laura", arrival_date="07/10/2026"),
    ]
    blocks, _ = build(guests, [])
    dates = [datetime.strptime(b.guest.arrival_date, "%m/%d/%Y") for b in blocks]
    assert dates == sorted(dates)
    assert blocks[0].guest.guest_name == "Hanks, Laura"


def test_guest_with_nothing_gets_single_blue_line():
    guest = make_guest(confirmation_number="C-1", guest_name="Swift, Michael")
    blocks, _ = build([guest], [])
    assert len(blocks[0].lines) == 1
    assert blocks[0].lines[0].reservation is None
    assert blocks[0].lines[0].color == BLUE


def test_note_yelp_mismatch_is_red():
    guest = make_guest(
        confirmation_number="D-1", guest_name="Clarkson, Olivia",
        arrival_date="07/10/2026", departure_date="07/12/2026",
        comments_notes="artisans on 7/10 @ 6pm",
    )
    reservation = make_reservation(
        guest_name="Olivia Clarkson", source_date="Jul 10, 2026",
        restaurant="Maggie's", notes_tags="maggie's",
    )
    blocks, _ = build([guest], [reservation])
    assert any(line.color == RED for line in blocks[0].lines)


def test_note_time_disagrees_with_yelp_time_is_red():
    """Same restaurant/date, but the note's time and Yelp's time disagree
    -- must render RED with the specific time-mismatch detail, not pass
    silently as a clean match."""
    guest = make_guest(
        confirmation_number="D2-1", guest_name="Smith, John",
        arrival_date="07/20/2026", departure_date="07/23/2026",
        comments_notes="Artisans 7/21 @ 7PM",
    )
    reservation = make_reservation(
        guest_name="John Smith", source_date="Jul 21, 2026", time="6:30", restaurant="Artisans",
    )
    blocks, _ = build([guest], [reservation])
    assert any(line.color == RED for line in blocks[0].lines)
    assert any("6:30" in line.note for line in blocks[0].lines)


def test_room_mismatch_is_yellow_not_red():
    guest = make_guest(
        confirmation_number="E-1", guest_name="Bieber, Daniel", room_name="LOOK",
        arrival_date="07/10/2026", departure_date="07/12/2026",
    )
    reservation = make_reservation(
        guest_name="Daniel Bieber", source_date="Jul 10, 2026",
        notes_tags="artisans KIWA", room_code_hint="s KIWA",
    )
    blocks, _ = build([guest], [reservation])
    assert any(line.color == YELLOW for line in blocks[0].lines)
    assert not any(line.color == RED for line in blocks[0].lines)


def test_full_name_no_room_hint_is_yellow():
    guest = make_guest(
        confirmation_number="E2-1", guest_name="Hopkins, Grace", room_name="AUSABL",
        arrival_date="07/10/2026", departure_date="07/13/2026",
    )
    reservation = make_reservation(
        guest_name="Grace Hopkins", source_date="Jul 11, 2026",
        notes_tags="artisans", room_code_hint=None,
    )
    blocks, _ = build([guest], [reservation])
    assert any(line.color == YELLOW for line in blocks[0].lines)
    assert any("No room found in the Yelp notes" in line.note for line in blocks[0].lines)


def test_surname_only_no_room_hint_is_red_potential_outside_guest():
    """The RED 'potential outside guest' case, per user direction: a bare
    surname match with no room hint is a real risk of confusing an
    unrelated person with an in-house guest, not just a soft data gap --
    must render RED, distinct from the softer YELLOW room/note issues."""
    guest = make_guest(
        confirmation_number="E3-1", guest_name="Hanks, Laura", room_name="TAHAW",
        arrival_date="07/10/2026", departure_date="07/13/2026",
    )
    reservation = make_reservation(
        guest_name="HANKS", source_date="Jul 11, 2026",
        notes_tags="artisans", room_code_hint=None,
    )
    blocks, _ = build([guest], [reservation])
    assert any(line.color == RED for line in blocks[0].lines)
    assert any("Potential outside guest" in line.note for line in blocks[0].lines)


def test_ambiguous_name_match_is_red_and_names_the_other_candidate():
    """Two unrelated same-surname guests, same night, no room hint or
    matching note time to break the tie -- the reservation must still
    show up (on the closer-scoring guest's card), flagged RED, naming who
    else it could belong to, rather than vanishing from the report."""
    guest_1 = make_guest(
        confirmation_number="H-1", guest_name="Depp, Emma", room_name="PINEVW",
        arrival_date="07/10/2026", departure_date="07/12/2026",
    )
    guest_2 = make_guest(
        confirmation_number="H-2", guest_name="Depp, Olivia", room_name="BUCK",
        arrival_date="07/09/2026", departure_date="07/15/2026",
    )
    reservation = make_reservation(
        guest_name="DEPP", source_date="Jul 10, 2026",
        notes_tags="artisans", room_code_hint=None,
    )
    blocks, _ = build([guest_1, guest_2], [reservation])
    matched_blocks = [
        b for b in blocks
        if any(line.reservation is reservation for line in b.lines)
    ]
    assert len(matched_blocks) == 1
    matched_line = next(line for line in matched_blocks[0].lines if line.reservation is reservation)
    assert matched_line.color == RED
    assert "Ambiguous match" in matched_line.note
    other_name = "Depp, Olivia" if matched_blocks[0].guest.confirmation_number == "H-1" else "Depp, Emma"
    assert other_name in matched_line.note


def test_ambiguous_name_match_resolved_by_note_time_is_yellow():
    """Same tie, but one candidate's notes mention this restaurant/date at
    the same time as the Yelp reservation -- that's enough to resolve it,
    shown as a softer YELLOW (inferred, not a clean name match) rather
    than the RED true-ambiguity case."""
    guest_1 = make_guest(
        confirmation_number="I-1", guest_name="Depp, Emma", room_name="PINEVW",
        arrival_date="07/10/2026", departure_date="07/12/2026",
        comments_notes="Artisans 7/10 @ 7PM",
    )
    guest_2 = make_guest(
        confirmation_number="I-2", guest_name="Depp, Olivia", room_name="BUCK",
        arrival_date="07/09/2026", departure_date="07/15/2026",
    )
    reservation = make_reservation(
        guest_name="DEPP", source_date="Jul 10, 2026", time="7:00",
        notes_tags="artisans", room_code_hint=None,
    )
    blocks, _ = build([guest_1, guest_2], [reservation])
    emma_block = next(b for b in blocks if b.guest.confirmation_number == "I-1")
    matched_line = next(line for line in emma_block.lines if line.reservation is reservation)
    assert matched_line.color == YELLOW
    assert "resolved using" in matched_line.note
    assert "Depp, Olivia" in matched_line.note


def test_reservation_without_note_is_yellow():
    guest = make_guest(confirmation_number="F-1", guest_name="Cyrus, Sophia")
    reservation = make_reservation(guest_name="Sophia Cyrus", source_date="Jul 10, 2026")
    blocks, _ = build([guest], [reservation])
    assert any(
        line.color == YELLOW and "not mentioned in Guest Notes" in line.note
        for line in blocks[0].lines
    )


def test_linked_group_note_on_every_block():
    guests = [
        make_guest(confirmation_number=f"G-{i}", guest_name="Jackson, John", room_name=room,
                   arrival_date="07/10/2026", departure_date="07/13/2026")
        for i, room in enumerate(["TAHAW", "AMPER", "MTJO"], start=1)
    ]
    blocks, _ = build(guests, [])
    jackson_blocks = by_name(blocks, "jackson")
    assert len(jackson_blocks) == 3
    for block in jackson_blocks:
        assert block.linked_group_note is not None


def test_room_move_note_present():
    earlier = make_guest(
        confirmation_number="H-1", guest_name="Sheeran, Steve",
        room_name="MOSS", room_type="1KF",
        arrival_date="07/10/2026", departure_date="07/12/2026",
    )
    later = make_guest(
        confirmation_number="H-2", guest_name="Sheeran, Steve",
        room_name="BIRCH", room_type="SBKP",
        arrival_date="07/12/2026", departure_date="07/14/2026",
    )
    blocks, room_moves = build([earlier, later], [])
    assert len(room_moves) == 1
    sheeran_blocks = by_name(blocks, "sheeran")
    assert any(b.room_move_note is not None for b in sheeran_blocks)


def test_no_line_left_without_a_color():
    guests = [
        make_guest(confirmation_number="I-1", guest_name="Depp, Emma"),
        make_guest(confirmation_number="I-2", guest_name="Hanks, Laura"),
    ]
    reservations = [make_reservation(guest_name="Emma Depp", source_date="Jul 10, 2026")]
    blocks, _ = build(guests, reservations)
    valid_colors = {GREEN, YELLOW, RED, BLUE}
    for block in blocks:
        for line in block.lines:
            assert line.color in valid_colors


def test_virtuoso_and_breakfast_flagged():
    guest = make_guest(
        confirmation_number="J-1", guest_name="Swift, Michael",
        reservation_notes="BREAKFAST INCLUDED VIRTUOSO...$100.00 VIRTUOSO CREDIT...",
    )
    blocks, _ = build([guest], [])
    assert blocks[0].virtuoso
    assert blocks[0].breakfast_included


def test_breakfast_typo_tolerated():
    guest = make_guest(
        confirmation_number="K-1", guest_name="Clarkson, Olivia",
        reservation_notes="BREAKFAST INCLLUDED as part of package.",
    )
    blocks, _ = build([guest], [])
    assert blocks[0].breakfast_included


def test_breakfast_without_virtuoso_does_not_set_virtuoso():
    guest = make_guest(
        confirmation_number="L-1", guest_name="Bieber, Daniel",
        guest_notes="Breakfast included as part of package.",
    )
    blocks, _ = build([guest], [])
    assert blocks[0].breakfast_included
    assert not blocks[0].virtuoso


def test_write_report_produces_valid_workbook(tmp_path):
    guest = make_guest(confirmation_number="M-1", guest_name="Cyrus, Sophia")
    blocks, _ = build([guest], [])

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out))
    assert out.exists()

    workbook = openpyxl.load_workbook(str(out))
    sheet = workbook.active
    assert sheet.max_row > len(blocks)


def test_room_move_count_appears_in_summary(tmp_path):
    earlier = make_guest(
        confirmation_number="N-1", guest_name="Jackson, John",
        room_name="MOSS", room_type="1KF",
        arrival_date="07/10/2026", departure_date="07/12/2026",
    )
    later = make_guest(
        confirmation_number="N-2", guest_name="Jackson, John",
        room_name="BIRCH", room_type="SBKP",
        arrival_date="07/12/2026", departure_date="07/14/2026",
    )
    blocks, room_moves = build([earlier, later], [])
    assert len(room_moves) == 1

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out), room_move_count=len(room_moves))

    workbook = openpyxl.load_workbook(str(out))
    summary_text = workbook.active.cell(row=2, column=1).value
    assert "1 room move" in summary_text


def test_group_booking_count_appears_in_summary(tmp_path):
    guests = [
        make_guest(confirmation_number=f"O-{i}", guest_name="Hanks, Laura", room_name=room,
                   arrival_date="07/10/2026", departure_date="07/13/2026")
        for i, room in enumerate(["TAHAW", "AMPER"], start=1)
    ]
    blocks, _ = build(guests, [])
    room_groups = find_room_groups(guests)
    assert len(room_groups) == 1

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out), group_booking_count=len(room_groups))

    workbook = openpyxl.load_workbook(str(out))
    summary_text = workbook.active.cell(row=2, column=1).value
    assert "1 group booking" in summary_text


def test_write_report_includes_rate_plan(tmp_path):
    guest = make_guest(confirmation_number="P-1", guest_name="Depp, Emma", rate_plan="LEIS")
    blocks, _ = build([guest], [])

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out))

    workbook = openpyxl.load_workbook(str(out))
    sheet = workbook.active
    labels = [cell.value for row in sheet.iter_rows(min_col=1, max_col=1) for cell in row]
    assert "Rate Plan" in labels
    values = [cell.value for row in sheet.iter_rows(min_col=2, max_col=2) for cell in row]
    assert "LEIS" in values


def test_write_report_includes_perk_badges(tmp_path):
    guest = make_guest(
        confirmation_number="Q-1", guest_name="Swift, Michael",
        reservation_notes="BREAKFAST INCLUDED VIRTUOSO...$100.00 VIRTUOSO CREDIT...",
    )
    blocks, _ = build([guest], [])

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out))

    workbook = openpyxl.load_workbook(str(out))
    sheet = workbook.active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert any(v == "Virtuoso" for v in values)
    assert any(v == "Breakfast included" for v in values)

    fills = {
        cell.fill.start_color.rgb[-6:]
        for row in sheet.iter_rows()
        for cell in row
        if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb
    }
    assert ORANGE in fills


def _is_border_filled(sheet, row: int, border_col: int) -> bool:
    cell = sheet.cell(row=row, column=border_col)
    return bool(
        cell.fill and cell.fill.start_color and cell.fill.start_color.rgb
        and cell.fill.start_color.rgb[-6:] == BLOCK_BORDER_FILL
    )


def test_blank_row_between_blocks_has_no_border_fill(tmp_path):
    """Regression test for an off-by-one in _write_block's border-fill
    loop: it used to paint the block-border column one row too far, onto
    the blank separator row write_report leaves between guest blocks --
    every block in the sheet had a stray filled cell where a blank gap
    should be."""
    guests = [
        make_guest(confirmation_number="R-1", guest_name="Jackson, John"),
        make_guest(confirmation_number="R-2", guest_name="Sheeran, Steve"),
    ]
    blocks, _ = build(guests, [])

    out = tmp_path / "report.xlsx"
    write_report(blocks, str(out))

    workbook = openpyxl.load_workbook(str(out))
    sheet = workbook.active
    border_col = LEFT_COL_COUNT + 1

    second_block_title_row = next(
        row[0].row for row in sheet.iter_rows(min_col=1, max_col=1)
        if row[0].value == "Sheeran, Steve"
    )
    blank_separator_row = second_block_title_row - 1
    assert not _is_border_filled(sheet, blank_separator_row, border_col)
